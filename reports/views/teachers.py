# reports/views/teachers.py
# -*- coding: utf-8 -*-
from __future__ import annotations

from core.observability import report_degraded as _degraded

from ._helpers import *
from ._helpers import (
    _safe_next_url,
    _get_active_school, _user_manager_schools,
)
from ..permissions import effective_user_role_label, is_school_manager
from ..gender_labels import school_gender_labels
from ..lab_kinds import LabKind
from ..staff_assignments import (
    apply_staff_assignment,
    assignment_matches,
    get_assignment,
)


def _decorate_manage_teacher_rows(teachers, *, active_school: Optional[School]) -> None:
    """Attach display-only labels that come from memberships, not legacy Role."""
    for teacher in teachers:
        role_label = effective_user_role_label(teacher, active_school=active_school)
        role_kind = ""

        try:
            is_manager_here = bool(getattr(teacher, "is_school_manager_in_active_school", False))
            if active_school is None:
                is_manager_here = is_school_manager(teacher)
        except Exception:
            is_manager_here = False

        job_title = (getattr(teacher, "school_job_title", "") or "").strip()
        if is_manager_here:
            role_kind = "manager"
        elif job_title == SchoolMembership.JobTitle.ADMIN_STAFF:
            role_kind = "admin_staff"
        elif job_title == SchoolMembership.JobTitle.LAB_TECH:
            role_kind = "lab_tech"
        elif bool(getattr(teacher, "has_teacher_membership", False)):
            role_kind = "teacher"
        elif role_label and role_label != "مستخدم":
            role_kind = "other"

        teacher.manage_role_kind = role_kind
        teacher.manage_role_label = "" if role_label == "مستخدم" else role_label

# =========================
# إدارة المعلّمين (مدير فقط)
# =========================
@login_required(login_url="reports:login")
@role_required({"manager"})  # إن كنت تبغى السماح للسوبر دائمًا، خلي role_required يتجاوز للسوبر أو أضف دور admin
@require_http_methods(["GET"])
def manage_teachers(request: HttpRequest) -> HttpResponse:
    active_school = _get_active_school(request)

    # ✅ اجبار اختيار مدرسة لغير السوبر (أوضح وأأمن)
    if not request.user.is_superuser:
        if active_school is None:
            messages.error(request, "فضلاً اختر مدرسة أولاً.")
            return redirect("reports:select_school")

        if active_school not in _user_manager_schools(request.user):
            messages.error(request, "ليست لديك صلاحية على هذه المدرسة.")
            return redirect("reports:select_school")

    term = (request.GET.get("q") or "").strip()
    status_filter = (request.GET.get("status") or "").strip()
    job_title_filter = (request.GET.get("job_title") or "").strip()
    department_filter = (request.GET.get("department") or "").strip()

    qs = Teacher.objects.order_by("-id")

    # ✅ عزل حسب المدرسة (نُظهر المعلمين المرتبطين بالمدرسة)
    if active_school is not None:
        qs = qs.filter(
            school_memberships__school=active_school,
            school_memberships__role_type__in=SchoolMembership.STAFF_ROLES,
        ).distinct()

    # ✅ بحث
    if term:
        qs = qs.filter(
            Q(name__icontains=term) |
            Q(phone__icontains=term) |
            Q(national_id__icontains=term)
        )
    if status_filter == "active":
        qs = qs.filter(is_active=True)
    elif status_filter == "inactive":
        qs = qs.filter(is_active=False)
    if active_school is not None and job_title_filter in SchoolMembership.JobTitle.values:
        qs = qs.filter(
            school_memberships__school=active_school,
            school_memberships__role_type__in=SchoolMembership.STAFF_ROLES,
            school_memberships__job_title=job_title_filter,
        )
    if active_school is not None and department_filter.isdigit():
        qs = qs.filter(
            dept_memberships__department_id=int(department_filter),
            dept_memberships__department__school=active_school,
        )
    qs = qs.distinct()

    # ✅ تمييز العضوية الحالية داخل المدرسة النشطة
    if active_school is not None:
        try:
            teacher_m = SchoolMembership.objects.filter(
                school=active_school,
                teacher=OuterRef("pk"),
                role_type__in=SchoolMembership.STAFF_ROLES,
                is_active=True,
            )
            title_sq = (
                teacher_m
                .values("job_title")[:1]
            )
            manager_m = SchoolMembership.objects.filter(
                school=active_school,
                teacher=OuterRef("pk"),
                role_type=SchoolMembership.RoleType.MANAGER,
                is_active=True,
            )
            qs = qs.annotate(
                has_teacher_membership=Exists(teacher_m),
                is_school_manager_in_active_school=Exists(manager_m),
                school_job_title=Subquery(title_sq),
            )
        except Exception:
            # بلا هذه التعليقات يفقد الكشف عمودَ الدور والمسمّى الوظيفي —
            # ويقرأ المدير كشفاً ناقصاً بلا ما يدلّ على نقصه.
            _degraded("teachers.annotate_membership_columns")

    # ✅ منع N+1: Prefetch عضويات الأقسام مرة واحدة وبحقول أقل
    if DepartmentMembership is not None:
        dm_qs = (
            DepartmentMembership.objects
            .select_related("department")
            .only("id", "teacher_id", "role_type", "department__id", "department__name", "department__slug")
            .order_by("department__name")
        )
        if active_school is not None:
            dm_qs = dm_qs.filter(Q(department__school=active_school) | Q(department__school__isnull=True))

        qs = qs.prefetch_related(Prefetch("dept_memberships", queryset=dm_qs))

    page = Paginator(qs, 20).get_page(request.GET.get("page"))
    _decorate_manage_teacher_rows(page.object_list, active_school=active_school)
    departments = (
        Department.objects.filter(school=active_school, is_active=True).order_by("name", "id")
        if active_school is not None
        else Department.objects.none()
    )
    current_count = SchoolMembership.seats_used(active_school)
    active_subscription = getattr(active_school, "subscription", None)
    maximum = int(getattr(active_subscription, "teacher_limit", 0) or 0)
    filters_query = request.GET.copy()
    filters_query.pop("page", None)
    return render(
        request,
        "reports/manage_teachers.html",
        {
            "teachers_page": page,
            "term": term,
            "status_filter": status_filter,
            "job_title_filter": job_title_filter,
            "department_filter": department_filter,
            "departments": departments,
            "job_title_choices": SchoolMembership.JobTitle.choices,
            "capacity": {
                "current": current_count,
                "maximum": maximum,
                "remaining": max(maximum - current_count, 0) if maximum else None,
            },
            "filters_query": filters_query.urlencode(),
        },
    )

@login_required(login_url="reports:login")
@role_required({"manager"})
@ratelimit(key="user", rate="30/h", method="POST", block=True)
@require_http_methods(["GET", "POST"])
def teacher_onboarding(request: HttpRequest) -> HttpResponse:
    """Unified, preview-first onboarding for one or many school users."""
    from ..teacher_onboarding import (
        RESULT_SESSION_KEY,
        build_preview,
        clear_preview,
        confirm_preview,
        load_preview,
        load_result,
        rows_from_quick_post,
        rows_from_uploaded_file,
        save_preview,
    )
    from ..staff_assignments import assignment_choices

    active_school = _get_active_school(request)
    if active_school is None:
        messages.error(request, "فضلاً اختر مدرسة أولاً.")
        return redirect("reports:select_school")
    if (not request.user.is_superuser) and active_school not in _user_manager_schools(request.user):
        messages.error(request, "ليست لديك صلاحية على هذه المدرسة.")
        return redirect("reports:select_school")

    subscription = getattr(active_school, "subscription", None)
    if subscription is None or bool(getattr(subscription, "is_expired", True)):
        messages.error(request, "لا يوجد اشتراك فعّال لهذه المدرسة.")
        return redirect("reports:my_subscription")

    def posted_quick_rows() -> list[dict[str, str]]:
        fields = {
            "name": request.POST.getlist("name"),
            "phone": request.POST.getlist("phone"),
            "national_id": request.POST.getlist("national_id"),
            "job_title": request.POST.getlist("job_title"),
            "department_id": request.POST.getlist("department"),
            "lab_kind": request.POST.getlist("lab_kind"),
        }
        count = max((len(values) for values in fields.values()), default=0)
        return [
            {
                field: values[index] if index < len(values) else ""
                for field, values in fields.items()
            }
            for index in range(count)
        ]

    quick_rows: list[dict[str, Any]] | None = None
    if request.method == "POST":
        action = (request.POST.get("action") or "").strip()
        remove_row = (request.POST.get("remove_row") or "").strip()

        if action == "add_quick_row":
            quick_rows = posted_quick_rows()
            if len(quick_rows) >= 2000:
                messages.error(request, "وصلت إلى الحد الأقصى وهو 2000 صف.")
            else:
                quick_rows.append(
                    {
                        "name": "",
                        "phone": "",
                        "national_id": "",
                        "job_title": SchoolMembership.JobTitle.TEACHER,
                        "department_id": "",
                        "lab_kind": "",
                    }
                )

        elif remove_row:
            quick_rows = posted_quick_rows()
            try:
                remove_index = int(remove_row)
            except (TypeError, ValueError):
                remove_index = -1
            if len(quick_rows) > 1 and 0 <= remove_index < len(quick_rows):
                quick_rows.pop(remove_index)
            elif quick_rows:
                quick_rows[0] = {
                    "name": "",
                    "phone": "",
                    "national_id": "",
                    "job_title": SchoolMembership.JobTitle.TEACHER,
                    "department_id": "",
                    "lab_kind": "",
                }

        elif action == "cancel":
            clear_preview(request)
            request.session.pop(RESULT_SESSION_KEY, None)
            messages.info(request, "تم إلغاء العملية دون حفظ أي بيانات.")
            return redirect("reports:bulk_import_teachers")

        elif action == "quick_preview":
            try:
                rows = rows_from_quick_post(request.POST)
                if not rows:
                    messages.error(request, "أدخل مستخدمًا واحدًا على الأقل قبل المتابعة.")
                else:
                    preview = build_preview(rows, active_school)
                    save_preview(request, preview, active_school, source="quick")
                    return redirect(f"{reverse('reports:bulk_import_teachers')}?step=preview")
            except ValueError as exc:
                if str(exc) == "too_many_rows":
                    messages.error(request, "القائمة تتجاوز الحد الأقصى وهو 2000 مستخدم.")
                else:
                    messages.error(request, "تعذر قراءة البيانات المدخلة. راجعها ثم حاول مجددًا.")

        elif action == "file_preview":
            uploaded_file = request.FILES.get("excel_file")
            if uploaded_file is None:
                messages.error(request, "اختر ملف Excel أو CSV أولاً.")
            else:
                try:
                    rows = rows_from_uploaded_file(uploaded_file)
                    if not rows:
                        messages.error(request, "الملف لا يحتوي على صفوف بيانات.")
                    else:
                        preview = build_preview(rows, active_school)
                        save_preview(request, preview, active_school, source="file")
                        return redirect(f"{reverse('reports:bulk_import_teachers')}?step=preview")
                except ValueError as exc:
                    errors = {
                        "unsupported_file": "صيغة الملف غير مدعومة. استخدم Excel أو CSV.",
                        "required_headers_missing": "يجب أن يحتوي الملف على عمودي الاسم الكامل ورقم الجوال.",
                        "too_many_rows": "الملف يتجاوز الحد الأقصى وهو 2000 صف.",
                        "file_too_large": "حجم الملف يتجاوز 10 ميجابايت. قسّمه إلى ملفات أصغر ثم حاول مجددًا.",
                    }
                    messages.error(request, errors.get(str(exc), "تعذر قراءة الملف. تأكد من سلامته ثم حاول مجددًا."))
                except Exception:
                    logger.exception("Teacher onboarding preview failed")
                    messages.error(request, "تعذر قراءة الملف. تأكد من سلامته ثم حاول مجددًا.")

        elif action == "confirm":
            try:
                result = confirm_preview(
                    request,
                    active_school,
                    request.POST.get("preview_token") or "",
                )
                messages.success(
                    request,
                    f"اكتملت العملية بنجاح لـ {result['total']} مستخدم.",
                )
                return redirect(f"{reverse('reports:bulk_import_teachers')}?completed=1")
            except ValueError as exc:
                if str(exc) == "preview_changed":
                    messages.warning(request, "تغيرت بعض البيانات منذ المعاينة. راجع النتائج المحدثة ثم أكد مرة أخرى.")
                    return redirect(f"{reverse('reports:bulk_import_teachers')}?step=preview")
                messages.error(request, "انتهت صلاحية المعاينة. أعد إدخال البيانات أو رفع الملف.")
            except Exception:
                logger.exception("Teacher onboarding confirmation failed")
                messages.error(request, "تعذر إتمام العملية ولم تُحفظ بيانات جزئية. حاول مرة أخرى.")

    preview = load_preview(request, active_school)
    result = load_result(request, active_school)
    current_count = SchoolMembership.seats_used(active_school)
    maximum = int(getattr(subscription, "teacher_limit", 0) or 0)
    capacity = {
        "current": current_count,
        "maximum": maximum,
        "remaining": max(maximum - current_count, 0) if maximum else None,
    }
    departments = Department.objects.filter(
        school=active_school,
        is_active=True,
    ).order_by("name", "id")
    if quick_rows is None:
        if preview and preview.get("source") == "quick":
            quick_rows = list(preview.get("rows") or [])
        else:
            quick_rows = [
                {
                    "name": "",
                    "phone": "",
                    "national_id": "",
                    "job_title": SchoolMembership.JobTitle.TEACHER,
                    "department_id": "",
                    "lab_kind": "",
                }
                for _ in range(3)
            ]
    return render(
        request,
        "reports/bulk_import_teachers.html",
        {
            "preview": preview,
            "result": result,
            "capacity": capacity,
            "departments": departments,
            "lab_kind_choices": LabKind.choices,
            "job_title_choices": assignment_choices(active_school),
            "active_mode": (request.GET.get("mode") or "quick").strip(),
            "quick_rows": quick_rows,
        },
    )


@login_required(login_url="reports:login")
@role_required({"manager"})
@require_http_methods(["GET"])
def bulk_import_teachers_issues(request: HttpRequest) -> HttpResponse:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from ..teacher_onboarding import load_preview

    active_school = _get_active_school(request)
    if active_school is None:
        return redirect("reports:select_school")
    preview = load_preview(request, active_school)
    if preview is None:
        messages.error(request, "لا توجد معاينة متاحة لتنزيل ملاحظاتها.")
        return redirect("reports:bulk_import_teachers")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ملاحظات الاستيراد"
    sheet.sheet_view.rightToLeft = True
    headers = ["الصف", "الاسم", "رقم الجوال", "الحالة", "الأخطاء", "التنبيهات"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="006C35")
        cell.alignment = Alignment(horizontal="center")
    for row in preview.get("rows") or []:
        if row.get("errors") or row.get("warnings"):
            sheet.append(
                [
                    row.get("row_number"),
                    row.get("name"),
                    row.get("phone"),
                    row.get("state_label"),
                    " | ".join(row.get("errors") or []),
                    " | ".join(row.get("warnings") or []),
                ]
            )
    for column, width in {"A": 10, "B": 28, "C": 18, "D": 22, "E": 55, "F": 55}.items():
        sheet.column_dimensions[column].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="teacher-import-issues.xlsx"'
    return response


@login_required(login_url="reports:login")
@role_required({"manager"})
@require_http_methods(["GET"])
def bulk_import_teachers_result(request: HttpRequest) -> HttpResponse:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from ..teacher_onboarding import load_result

    active_school = _get_active_school(request)
    if active_school is None:
        return redirect("reports:select_school")
    result = load_result(request, active_school)
    if result is None or not result.get("created_rows"):
        messages.error(request, "لا توجد حسابات جديدة لتنزيل بيانات دخولها.")
        return redirect("reports:bulk_import_teachers")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "بيانات الدخول"
    sheet.sheet_view.rightToLeft = True
    sheet.append(["الاسم", "رقم الجوال", "كلمة المرور المؤقتة", "تعليمات"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="006C35")
        cell.alignment = Alignment(horizontal="center")
    for row in result["created_rows"]:
        sheet.append(
            [
                row["name"],
                row["phone"],
                row["temporary_password"],
                "يجب تغيير كلمة المرور عند تسجيل الدخول لأول مرة.",
            ]
        )
    for column, width in {"A": 30, "B": 18, "C": 22, "D": 52}.items():
        sheet.column_dimensions[column].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="new-teachers-login.xlsx"'
    return response


@login_required(login_url="reports:login")
@role_required({"manager"})
@require_http_methods(["GET"])
def bulk_import_teachers_template(request: HttpRequest) -> HttpResponse:
    """Download a clean template; examples live on a separate instructions sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    active_school = _get_active_school(request)
    labels = school_gender_labels(active_school)
    from ..staff_assignments import assignment_choices

    job_labels = [label for _, label in assignment_choices(active_school)]
    sample_name = "نورة أحمد الغامدي" if labels["is_girls"] else "محمد أحمد الغامدي"
    wb = Workbook()
    ws = wb.active
    ws.title = str(labels["teachers"])
    ws.sheet_view.rightToLeft = True

    headers = ["الاسم الكامل", "رقم الجوال", "رقم الهوية", "المسمى الوظيفي", "القسم", "المختبر"]

    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill("solid", fgColor="006C35")
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center", readingOrder=2)
    thin = Side(style="thin", color="D6E4DC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, head in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=head)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 24
    ws.column_dimensions["F"].width = 24
    ws.freeze_panes = "A2"

    job_validation = DataValidation(
        type="list",
        formula1=f'"{",".join(job_labels)}"',
        allow_blank=True,
    )
    job_validation.error = "اختر مسمى وظيفيًا من القائمة."
    job_validation.errorTitle = "مسمى غير صحيح"
    ws.add_data_validation(job_validation)
    job_validation.add("D2:D2001")

    instructions = wb.create_sheet("التعليمات والأمثلة")
    instructions.sheet_view.rightToLeft = True
    instruction_rows = [
        ["الحقل", "هل هو مطلوب؟", "مثال", "ملاحظة"],
        ["الاسم الكامل", "مطلوب", sample_name, "يظهر بهذا الشكل داخل النظام."],
        ["رقم الجوال", "مطلوب", "0551234567", "اسم الدخول وكلمة المرور المؤقتة."],
        ["رقم الهوية", "اختياري", "1012345678", "10 أرقام عند إدخاله."],
        ["المسمى الوظيفي", "اختياري", labels["teacher_indefinite"], " أو ".join(job_labels) + "."],
        ["القسم", "اختياري", "قسم العلوم", "القسم التنظيمي المرتبط بالتقارير."],
        [
            "المختبر",
            "مطلوب لمحضر المختبر",
            "مختبر العلوم",
            "اختر مختبر العلوم أو مختبر الحاسب الآلي؛ ولا يرتبط هذا الحقل بأقسام التقارير.",
        ],
        ["تنبيه", "", "", f"لا تنسخ صفوف الأمثلة إلى ورقة {labels['teachers']}."],
    ]
    for row in instruction_rows:
        instructions.append(row)
    for cell in instructions[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
    for row in instructions.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = right
    for column, width in {"A": 24, "B": 17, "C": 28, "D": 55}.items():
        instructions.column_dimensions[column].width = width

    instructions.append([])
    if active_school is not None:
        department_names = list(
            Department.objects.filter(school=active_school, is_active=True)
            .order_by("name")
            .values_list("name", flat=True)
        )
        if department_names:
            instructions.append(["الأقسام المتاحة في المدرسة"])
            instructions.cell(row=instructions.max_row, column=1).font = Font(bold=True, color="006C35")
            for name in department_names:
                instructions.append([name])
            instructions.append([])
    instructions.append(["المختبرات المتاحة"])
    instructions.cell(row=instructions.max_row, column=1).font = Font(bold=True, color="006C35")
    for _value, label in LabKind.choices:
        instructions.append([label])

    from io import BytesIO
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="teachers-import-template.xlsx"'
    return response


@login_required(login_url="reports:login")
@role_required({"manager"})
@require_http_methods(["GET", "POST"])
def add_teacher(request: HttpRequest) -> HttpResponse:
    """إضافة منسوب وتكليفه من الكتالوج نفسه المستخدم في شاشة الأدوار."""
    active_school = _get_active_school(request)
    if School.objects.filter(is_active=True).exists():
        if active_school is None:
            messages.error(request, "فضلاً اختر مدرسة أولاً.")
            return redirect("reports:select_school")
        if (not request.user.is_superuser) and active_school not in _user_manager_schools(request.user):
            messages.error(request, "ليست لديك صلاحية على هذه المدرسة.")
            return redirect("reports:select_school")

    def render_form(form, *, confirmation=None):
        return render(
            request,
            "reports/add_teacher.html",
            {
                "form": form,
                "title": "إضافة منسوب",
                "page_title": "إضافة منسوب جديد",
                "page_subtitle": "أنشئ الحساب وحدّد دوره الأساسي في خطوة واضحة وآمنة",
                "role_change_confirmation": confirmation,
                "is_staff_onboarding": True,
                "next_url_value": _safe_next_url(
                    request.POST.get("next") or request.GET.get("next")
                ),
            },
        )

    if request.method != "POST":
        return render_form(TeacherCreateForm(active_school=active_school))

    continue_adding = request.POST.get("save_and_add_another") == "1"
    phone_raw = (request.POST.get("phone") or "").strip()
    existing_teacher = Teacher.objects.filter(phone=phone_raw).first() if phone_raw else None
    # تمرير instance يمنع تحقق التفرد من رفض الحساب الموجود؛ ولا يُحفظ هذا
    # النموذج عليه إطلاقاً، فتظل هويته وكلمة مروره كما هما.
    form = TeacherCreateForm(
        request.POST,
        active_school=active_school,
        instance=existing_teacher,
    )
    if not form.is_valid():
        messages.error(request, "الرجاء تصحيح الأخطاء الظاهرة.")
        return render_form(form)

    assignment_code = form.cleaned_data["job_title"]
    keep_teaching = bool(form.cleaned_data.get("keep_teaching_role"))
    selected_lab_kind = form.cleaned_data.get("lab_kind") or ""
    assignment = get_assignment(assignment_code)
    assignment_label = next(
        (
            item["label"]
            for item in form.assignment_cards
            if item["code"] == assignment_code
        ),
        assignment_code,
    )
    next_url = _safe_next_url(request.POST.get("next") or request.GET.get("next"))

    # الاشتراك والمقعد يُفحصان مرة واحدة بالطريقة نفسها للحساب الجديد والموجود.
    subscription = getattr(active_school, "subscription", None)
    if subscription is None or bool(getattr(subscription, "is_expired", True)):
        messages.error(request, "لا يوجد اشتراك فعّال لهذه المدرسة.")
        return render_form(form)
    maximum = int(getattr(subscription, "teacher_limit", 0) or 0)
    consumes_new_seat = not (
        existing_teacher
        and SchoolMembership.objects.filter(
            school=active_school,
            teacher=existing_teacher,
            role_type__in=SchoolMembership.SEAT_CONSUMING_ROLES,
        ).exists()
    )
    if (
        maximum > 0
        and consumes_new_seat
        and SchoolMembership.seats_used(active_school) >= maximum
    ):
        messages.error(
            request,
            f"لا يمكن إضافة أكثر من {maximum} من حسابات منسوبي المدرسة حسب الباقة.",
        )
        return render_form(form)

    try:
        if existing_teacher is not None:
            # أعد القراءة لأن ModelForm يبني القيم المنشورة على instance في الذاكرة؛
            # الربط لا يغيّر اسم الحساب الموجود أو هويته أو كلمة مروره.
            existing_teacher = Teacher.objects.get(pk=existing_teacher.pk)
            if not existing_teacher.is_active:
                messages.error(
                    request,
                    "الحساب الموجود موقوف على مستوى المنصة، ولا يمكن ربطه قبل إعادة تفعيله.",
                )
                return render_form(form)

            has_membership_here = SchoolMembership.objects.filter(
                school=active_school,
                teacher=existing_teacher,
                role_type__in=SchoolMembership.STAFF_ROLES,
            ).exists()
            exact_assignment = assignment_matches(
                school=active_school,
                member=existing_teacher,
                code=assignment_code,
                keep_teaching_role=keep_teaching,
            )
            if exact_assignment:
                if assignment_code == SchoolMembership.JobTitle.LAB_TECH:
                    SchoolMembership.objects.filter(
                        school=active_school,
                        teacher=existing_teacher,
                        role_type=assignment.role_type,
                    ).update(lab_kind=selected_lab_kind)
                messages.info(request, "الحساب مرتبط بالفعل بهذه المدرسة بالتكليف المختار.")
                return redirect(
                    "reports:add_teacher"
                    if continue_adding
                    else (next_url or "reports:manage_teachers")
                )

            if has_membership_here and request.POST.get("confirm_role_change") != "1":
                return render_form(
                    form,
                    confirmation={
                        "member_name": existing_teacher.name,
                        "current_role": effective_user_role_label(
                            existing_teacher, active_school=active_school
                        ),
                        "new_role": assignment_label,
                        "keeps_teaching": keep_teaching,
                    },
                )

            with transaction.atomic():
                membership = apply_staff_assignment(
                    school=active_school,
                    member=existing_teacher,
                    code=assignment_code,
                    keep_teaching_role=keep_teaching,
                    actor=request.user,
                )
                membership.lab_kind = selected_lab_kind
                membership.save(update_fields=["lab_kind"])
            messages.success(
                request,
                f"تم ربط الحساب الموجود وإسناد دور «{assignment_label}» دون تغيير بيانات دخوله.",
            )
        else:
            with transaction.atomic():
                teacher = form.save(commit=True)
                membership = apply_staff_assignment(
                    school=active_school,
                    member=teacher,
                    code=assignment_code,
                    keep_teaching_role=keep_teaching,
                    actor=request.user,
                )
                membership.lab_kind = selected_lab_kind
                membership.save(update_fields=["lab_kind"])
            messages.success(
                request,
                "تمت إضافة المنسوب. كلمة المرور المؤقتة هي رقم الجوال، وسيُطلب تغييرها عند أول دخول.",
            )

        if continue_adding:
            return redirect("reports:add_teacher")
        if next_url:
            return redirect(next_url)
        if assignment.requires_scope:
            messages.info(request, "أكمل الآن تحديد نطاق العمل والصلاحيات المناسبة.")
            return redirect("reports:staff_role_scope", pk=membership.pk)
        return redirect("reports:manage_teachers")
    except IntegrityError:
        messages.error(request, "تعذّر الحفظ: قد يكون رقم الجوال أو الهوية مستخدمًا مسبقًا.")
    except ValidationError as exc:
        messages.error(request, " ".join(getattr(exc, "messages", []) or [str(exc)]))
    except Exception:
        logger.exception("add_teacher failed")
        messages.error(request, "حدث خطأ غير متوقع أثناء الحفظ. جرّب لاحقًا.")
    return render_form(form)

@login_required(login_url="reports:login")
@role_required({"manager"})
@require_http_methods(["GET", "POST"])
def edit_teacher(request: HttpRequest, pk: int) -> HttpResponse:
    active_school = _get_active_school(request)
    teacher = get_object_or_404(Teacher, pk=pk)

    # لا يُسمح للمدير بتعديل من ليس من منسوبي مدرسته.
    #
    # **المنسوب لا المعلّم.** كان الشرط ``role_type=TEACHER`` وحده، وكشف
    # المنسوبين يقرأ ``STAFF_ROLES`` — فمن أُسند وكيلاً أو محضّر مختبر من شاشة
    # الأدوار (وهي تحذف عضويته التدريسية ما لم يُطلب الاحتفاظ بالنصاب) يظهر في
    # الكشف بزرَّي التعديل والحذف، ويردّه الزرّان بأنه «غير مرتبط بمدرستك».
    #
    # **بلا مدرسة نشطة يُمنع لا يُسمح.** كان الشرط ``and active_school is not
    # None``، فانعدامُها يُسقط فحص الارتباط كلَّه ويفتح أي حساب في المنصة برقمه.
    # وانعدامها وارد: ``role_required`` لا يفرض اختيار مدرسة إلا إن وُجدت مدرسة
    # واحدة مفعّلة على الأقل، فمدير مدرسةٍ عُطّلت يبلغ هذا الموضع بلا مدرسة.
    if not getattr(request.user, "is_superuser", False):
        if active_school is None:
            messages.error(request, "فضلاً اختر مدرسة أولاً.")
            return redirect("reports:select_school")
        has_membership = SchoolMembership.objects.filter(
            school=active_school,
            teacher=teacher,
            role_type__in=SchoolMembership.STAFF_ROLES,
            is_active=True,
        ).exists()
        if not has_membership:
            messages.error(request, "لا يمكنك تعديل بيانات هذا المستخدم لأنه ليس من منسوبي مدرستك.")
            return redirect("reports:manage_teachers")
    if request.method == "POST":
        # تعديل بيانات المعلّم فقط — التكاليف تتم من صفحة أعضاء القسم
        form = TeacherEditForm(request.POST, instance=teacher, active_school=active_school)
        if form.is_valid():
            try:
                with transaction.atomic():
                    form.save(commit=True)
                messages.success(request, "تم تحديث بيانات المستخدم بنجاح.")
                next_url = _safe_next_url(request.POST.get("next") or request.GET.get("next"))
                return redirect(next_url or "reports:manage_teachers")
            except Exception:
                logger.exception("edit_teacher failed")
                messages.error(request, "حدث خطأ غير متوقع أثناء التحديث.")
        else:
            messages.error(request, "الرجاء تصحيح الأخطاء الظاهرة.")
    else:
        form = TeacherEditForm(instance=teacher, active_school=active_school)

    return render(request, "reports/edit_teacher.html", {"form": form, "teacher": teacher, "title": "تعديل مستخدم"})

@login_required(login_url="reports:login")
@role_required({"manager"})
@require_http_methods(["POST"])
def delete_teacher(request: HttpRequest, pk: int) -> HttpResponse:
    active_school = _get_active_school(request)
    teacher = get_object_or_404(Teacher, pk=pk)

    # لا يُسمح للمدير بحذف من ليس من منسوبي مدرسته. والمدير نفسه خارج
    # ``STAFF_ROLES``، فلا يبلغ هذا الشرطَ من يحاول حذف مدير.
    #
    # **بلا مدرسة نشطة يُمنع لا يُسمح.** الخطر هنا أشدّ منه في التعديل: الفرع
    # الأخير أدناه ينفّذ ``teacher.delete()`` — حذفَ الحساب من المنصة كلها لا
    # فصلَه عن مدرسة. فبقاء الشرط مشروطاً بوجود المدرسة كان يجعل غيابَها طريقاً
    # إلى حذف أي حساب برقمه.
    if not getattr(request.user, "is_superuser", False):
        if active_school is None:
            messages.error(request, "فضلاً اختر مدرسة أولاً.")
            return redirect("reports:select_school")
        has_membership = SchoolMembership.objects.filter(
            school=active_school,
            teacher=teacher,
            role_type__in=SchoolMembership.STAFF_ROLES,
            is_active=True,
        ).exists()
        if not has_membership:
            messages.error(request, "لا يمكنك حذف هذا المستخدم لأنه ليس من منسوبي مدرستك.")
            return redirect("reports:manage_teachers")
    try:
        with transaction.atomic():
            # الحذف العالمي لمالك النظام وحده — وغيره لا يبلغ هنا إلا بمدرسة.
            if active_school is not None and not getattr(request.user, "is_superuser", False):
                # ✅ في وضع تعدد المدارس: لا نحذف الحساب عالميًا، بل نفصل عضويته عن هذه المدرسة فقط
                #
                # **كل أدواره لا دوره التدريسي.** حذف صفّ ``TEACHER`` وحده يترك
                # وكيلاً بعضوية وكالته: تقول الرسالة «أُزيل» وهو باقٍ بصلاحيته
                # وبمقعده. والنطاقات والتفويضات تسقط مع العضوية بـ CASCADE،
                # وهو الصحيح — نطاق من لم يعد منسوباً لا معنى له.
                SchoolMembership.objects.filter(
                    school=active_school,
                    teacher=teacher,
                    role_type__in=SchoolMembership.STAFF_ROLES,
                ).delete()
                messages.success(request, "تمت إزالة المستخدم من المدرسة الحالية.")
            else:
                teacher.delete()
                messages.success(request, "تم حذف المستخدم.")
    except Exception:
        logger.exception("delete_teacher failed")
        messages.error(request, "تعذّر حذف المستخدم. حاول لاحقًا.")
    next_url = _safe_next_url(request.POST.get("next") or request.GET.get("next"))
    return redirect(next_url or "reports:manage_teachers")
